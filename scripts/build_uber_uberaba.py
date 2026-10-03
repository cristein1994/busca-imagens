#!/usr/bin/env python3
"""Build Uber-in-Uberaba public research datasets."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path("/workspace/data")
OUT.mkdir(parents=True, exist_ok=True)

COMBINED_COLS = [
    "name",
    "title",
    "type",
    "address",
    "location",
    "phone",
    "website",
    "linkedinUrl",
    "email",
    "notes",
]

people = [
    {
        "name": "Gabriel Soares da Penha",
        "title": "Motorista autônomo (Aplicativos) - Uber",
        "type": "person_driver",
        "address": "",
        "location": "Uberaba, Minas Gerais, Brazil",
        "phone": "",
        "website": "",
        "linkedinUrl": "https://www.linkedin.com/in/gabriel-soares-da-penha-86b2362bb",
        "email": "",
        "notes": "Public LinkedIn; current Uber partner driver role since Jan 2023 in Uberaba.",
    },
    {
        "name": "Rodrigo Junqueira Souto",
        "title": "Motorista de aplicativo - Uber e 99",
        "type": "person_driver",
        "address": "",
        "location": "Uberaba, Minas Gerais, Brazil",
        "phone": "",
        "website": "",
        "linkedinUrl": "https://www.linkedin.com/in/rodrigo-junqueira-souto-2a164810b",
        "email": "",
        "notes": "Public LinkedIn; rideshare driver Uber+99 since May 2021 in Uberaba.",
    },
    {
        "name": "Deivy Alexandre Goes",
        "title": "Motorista - Uber",
        "type": "person_driver",
        "address": "",
        "location": "Uberaba, Minas Gerais, Brazil",
        "phone": "",
        "website": "",
        "linkedinUrl": "https://www.linkedin.com/in/deivy-alexandre-goes-a0ab9410a",
        "email": "",
        "notes": "Public LinkedIn headline 'Motorista da empresa Uber'; current role since Jul 2024 in Uberaba.",
    },
    {
        "name": "PAULO MIGUEZ",
        "title": "Motorista de aplicativos - Uber",
        "type": "person_driver",
        "address": "",
        "location": "Uberaba, Minas Gerais, Brazil",
        "phone": "",
        "website": "",
        "linkedinUrl": "https://www.linkedin.com/in/paulo-miguez-5a9936213",
        "email": "",
        "notes": "Public LinkedIn; current Uber partner driver role since Dec 2022 in Uberaba.",
    },
    {
        "name": "Eurípedes Junior",
        "title": "Motorista de automóvel - Uber",
        "type": "person_driver",
        "address": "",
        "location": "Uberaba, Minas Gerais, Brazil",
        "phone": "",
        "website": "",
        "linkedinUrl": "https://www.linkedin.com/in/eur%C3%ADpedes-junior-405441244",
        "email": "",
        "notes": "Public LinkedIn; current Uber driver role since Feb 2025 in Uberaba.",
    },
    {
        "name": "Luiz Filipe Lima Martins",
        "title": "Motorista - Uber",
        "type": "person_driver",
        "address": "",
        "location": "Uberaba, Minas Gerais, Brazil",
        "phone": "",
        "website": "",
        "linkedinUrl": "https://www.linkedin.com/in/luiz-filipe-lima-martins-152226371",
        "email": "",
        "notes": "Public LinkedIn; current Uber driver role since Mar 2022 in Uberaba.",
    },
]

places = [
    {
        "name": "Uber — cidade Uberaba",
        "title": "Página oficial da cidade (passageiros)",
        "type": "uber_city_page",
        "address": "Uberaba, MG, Brasil",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://www.uber.com/global/pt-br/r/cities/uberaba-mg-br/",
        "linkedinUrl": "",
        "email": "",
        "notes": "Confirma app disponível 24/7; destinos populares (Shopping, ABCZ, Praça Rui Barbosa); menciona Uber Eats.",
    },
    {
        "name": "Uber — cadastro motorista Uberaba",
        "title": "Página de signup para motoristas parceiros",
        "type": "driver_signup",
        "address": "Uberaba, MG, Brasil",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://www.uber.com/br/pt-br/e/drive/uberaba-mg-br/",
        "linkedinUrl": "",
        "email": "",
        "notes": "Cadastro aberto para motoristas; também promove entregas Uber Eats (2 ou 4 rodas). Requisitos: CNH, veículo elegível, checagem de segurança, 1+ ano experiência.",
    },
    {
        "name": "Uber Envios — Uberaba",
        "title": "Courier / entregas locais",
        "type": "uber_courier_page",
        "address": "Uberaba, MG, Brasil",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://www.uber.com/br/pt-br/b/courier-services/uberaba-mg-br/",
        "linkedinUrl": "",
        "email": "",
        "notes": "Uber Envios / Uber Direct / opções moto e carro anunciadas para Uberaba.",
    },
    {
        "name": "Uber — Aeroporto de Uberaba (UBA)",
        "title": "Página aeroporto / pickup",
        "type": "uber_airport_page",
        "address": "Aeroporto Mário de Almeida Franco (UBA), Uberaba, MG",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://www.uber.com/global/pt-br/r/airports/uba/?id=542",
        "linkedinUrl": "",
        "email": "",
        "notes": "Lista opções UberX, Priority, Uber Mulher, Moto; preços médios ilustrativos. Página oficial mistura texto de endereço com Ubá em um trecho — tratar UBA como Aeroporto de Uberaba.",
    },
    {
        "name": "UBER DO BRASIL TECNOLOGIA LTDA. (filial Uberaba)",
        "title": "Filial CNPJ baixada",
        "type": "corporate_branch_closed",
        "address": "Avenida Guilherme Ferreira, 1760 - São Benedito, Uberaba - MG, 38022-200",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://empresadois.com.br/cnpj/uber-do-brasil-tecnologia-ltda-17895646007270",
        "linkedinUrl": "",
        "email": "",
        "notes": "CNPJ 17.895.646/0072-70; abertura 27/11/2019; situação BAIXADA (não é hub ativo). CNAE intermediação/agenciamento. Fonte: registros públicos CNPJ (EmpresaDois/Econodata).",
    },
    {
        "name": "Uber Greenlight / Espaço Uber Uberaba",
        "title": "Hub presencial parceiros",
        "type": "greenlight_hub",
        "address": "",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://www.uber.com/br/pt-br/drive/contact/",
        "linkedinUrl": "",
        "email": "",
        "notes": "Nenhum Espaço Uber / Greenlight público listado para Uberaba. Suporte presencial em grandes cidades (ex.: SP Barra Funda); checar app parceiro > Ajuda > Agendar atendimento presencial.",
    },
    {
        "name": "G1 — Uber inicia operações em Uberaba",
        "title": "Notícia de lançamento (14/12/2016)",
        "type": "news",
        "address": "Uberaba, MG",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://g1.globo.com/minas-gerais/triangulo-mineiro/noticia/2016/12/uber-inicia-oficialmente-operacoes-em-uberaba.html",
        "linkedinUrl": "",
        "email": "",
        "notes": "Lançamento oficial 14/12/2016 às 14h; modalidade inicial mais barata; pagamento só cartão no início. Sedest alertou sobre legislação local.",
    },
    {
        "name": "Diário do Transporte — regulamentação apps Uberaba",
        "title": "Notícia regulamentação municipal (2018)",
        "type": "news",
        "address": "Uberaba, MG",
        "location": "Uberaba, MG",
        "phone": "",
        "website": "https://diariodotransporte.com.br/2018/11/20/transporte-por-aplicativo-e-regulamentado-em-uberaba/",
        "linkedinUrl": "",
        "email": "",
        "notes": "Lei municipal 12.963/2018; Prefeitura citou ~1.500 motoristas cadastrados no Uber à época; regras para Uber/99.",
    },
    {
        "name": "99 — cidade Uberaba (contexto concorrente)",
        "title": "Cadastro motorista 99 em Uberaba",
        "type": "competitor_city_page",
        "address": "Uberaba, MG",
        "location": "Uberaba, MG",
        "phone": "0300 313 2421",
        "website": "https://99app.com/cidades/mg/uberaba/",
        "linkedinUrl": "",
        "email": "",
        "notes": "99 opera em Uberaba (99Pop/99Táxi). Casa99 presencial mais próxima listada em BH (não há hub 99 em Uberaba).",
    },
]

runs = {
    "apify": {
        "status": "blocked_monthly_limit",
        "attempted": [
            {
                "actor": "apify/rag-web-browser",
                "purpose": "Web search Uber Uberaba / Uber Eats",
                "result": "Monthly usage hard limit exceeded",
            },
            {
                "actor": "harvestapi/linkedin-profile-search",
                "purpose": "LinkedIn people Uber in Uberaba (Short mode)",
                "input": {
                    "profileScraperMode": "Short",
                    "searchQuery": "Uber",
                    "locations": ["Uberaba"],
                    "maxItems": 25,
                },
                "result": "Monthly usage hard limit exceeded",
            },
        ],
    },
    "fallback": {
        "status": "ok",
        "sources": [
            "Cursor WebSearch / WebFetch",
            "Official Uber city / drive / courier / airport pages",
            "Public LinkedIn profile snippets (Google/Bing indexed)",
            "Public CNPJ directories (EmpresaDois / Econodata)",
            "G1 Triângulo Mineiro; Diário do Transporte",
            "99 city page (competitor context only)",
        ],
    },
    "collectedAt": "2026-10-03",
    "geo": "Uberaba, Minas Gerais, Brazil",
}

summary = {
    "title": "Uber em Uberaba / MG — presença pública",
    "status": {
        "rides_active": True,
        "uber_eats_promoted": True,
        "uber_envios_promoted": True,
        "greenlight_or_espaco_uber_local": False,
        "cnpj_branch": "baixada",
        "launch_date_reported": "2016-12-14",
        "municipal_regulation": "Lei 12.963/2018",
        "drivers_reported_by_prefeitura_2018": 1500,
    },
    "counts": {
        "people": len(people),
        "places_pages": len(places),
        "combined": len(people) + len(places),
        "people_with_linkedin": sum(1 for p in people if p["linkedinUrl"]),
        "places_with_website": sum(1 for p in places if p["website"]),
        "corporate_staff_uberaba": 0,
        "greenlight_hubs_uberaba": 0,
    },
    "apifyStatus": "blocked_monthly_limit",
    "primarySources": runs["fallback"]["sources"],
    "keyFindings": [
        "Uber rides operate in Uberaba (official city + driver pages; airport UBA).",
        "Uber Eats and Uber Envios are promoted for Uberaba on official Uber pages.",
        "No public Espaço Uber / Greenlight hub found in Uberaba.",
        "Uber Brasil CNPJ branch 17.895.646/0072-70 at Av. Guilherme Ferreira 1760 is BAIXADA.",
        "6 public LinkedIn partner-driver profiles located in Uberaba; no local corporate Uber staff found.",
        "99 also has a city driver page; nearest Casa99 listed in Belo Horizonte.",
    ],
}


def write_json(path: Path, data) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict], fieldnames: list[str], *, excel_ios: bool = False) -> None:
    if excel_ios:
        # UTF-8 BOM + CRLF + semicolon for Excel/iOS Numbers
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(
                f,
                fieldnames=fieldnames,
                delimiter=";",
                lineterminator="\r\n",
                extrasaction="ignore",
            )
            writer.writeheader()
            for row in rows:
                writer.writerow({k: row.get(k, "") for k in fieldnames})
    else:
        with path.open("w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({k: row.get(k, "") for k in fieldnames})


def write_xlsx(path: Path, sheets: dict[str, list[dict]], cols: list[str]) -> None:
    wb = Workbook()
    first = True
    for title, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(title[:31])
        if first:
            ws.title = title[:31]
            first = False
        for col, name in enumerate(cols, 1):
            cell = ws.cell(1, col, name)
            cell.font = Font(bold=True)
        for r_i, row in enumerate(rows, 2):
            for c_i, name in enumerate(cols, 1):
                ws.cell(r_i, c_i, row.get(name, ""))
        for col in ws.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col:
                max_len = max(max_len, len(str(cell.value or "")))
            ws.column_dimensions[letter].width = min(max_len + 2, 60)
    wb.save(path)


def write_html(path: Path, combined: list[dict]) -> None:
    rows_html = []
    for r in combined:
        cells = "".join(f"<td>{(r.get(c) or '').replace('<','&lt;')}</td>" for c in COMBINED_COLS)
        rows_html.append(f"<tr>{cells}</tr>")
    html = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8"/>
<title>Uber em Uberaba — dados públicos</title>
<style>
body{{font-family:Georgia,serif;margin:2rem;background:#f7f4ef;color:#1a1a1a}}
h1{{font-size:1.6rem}}
table{{border-collapse:collapse;width:100%;background:#fff}}
th,td{{border:1px solid #ddd;padding:.45rem .55rem;vertical-align:top;font-size:.85rem}}
th{{background:#0e3b2e;color:#fff;text-align:left}}
tr:nth-child(even){{background:#f3f8f5}}
.note{{color:#444;max-width:70ch}}
</style>
</head>
<body>
<h1>Uber em Uberaba / MG</h1>
<p class="note">Dados públicos agregados (páginas Uber, notícias, CNPJ baixado, LinkedIn indexado). Sem números privados.</p>
<table>
<thead><tr>{''.join(f'<th>{c}</th>' for c in COMBINED_COLS)}</tr></thead>
<tbody>
{''.join(rows_html)}
</tbody>
</table>
</body>
</html>
"""
    path.write_text(html, encoding="utf-8")


def main() -> None:
    combined = people + places

    write_json(OUT / "uber-uberaba-people.json", people)
    write_json(OUT / "uber-uberaba-places.json", places)
    write_json(OUT / "uber-uberaba-runs.json", runs)
    write_json(OUT / "uber-uberaba-summary.json", summary)
    # aliases requested in task
    write_json(OUT / "runs.json", runs)
    write_json(OUT / "summary.json", summary)

    people_cols = COMBINED_COLS
    places_cols = COMBINED_COLS

    write_csv(OUT / "uber-uberaba-people.csv", people, people_cols)
    write_csv(OUT / "uber-uberaba-places.csv", places, places_cols)
    write_csv(OUT / "uber-uberaba-combined.csv", combined, COMBINED_COLS, excel_ios=True)
    write_csv(OUT / "uber-uberaba-ios.csv", combined, COMBINED_COLS, excel_ios=True)

    write_xlsx(
        OUT / "uber-uberaba.xlsx",
        {"combined": combined, "people": people, "places": places},
        COMBINED_COLS,
    )
    write_html(OUT / "uber-uberaba.html", combined)

    print(f"Wrote {len(people)} people, {len(places)} places -> {OUT}")


if __name__ == "__main__":
    main()
